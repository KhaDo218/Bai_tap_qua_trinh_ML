import html
import os

import pandas as pd
import plotly.express as px
import plotly.graph_objects as go
from plotly.offline import get_plotlyjs
from openpyxl.styles import Alignment, Font, PatternFill

CLUSTER_COLORS = ["#4361ee", "#f72585", "#2a9d8f", "#fb8500", "#7209b7", "#e63946", "#8d99ae", "#118ab2"]


# ======================================================================
# 1. XUẤT EXCEL: mỗi cụm 1 sheet + sheet "Thống kê" (giống bài thầy)
#    df phải có cột "cluster" (đánh số từ 1) và "ClusterName"
# ======================================================================
def exportClustersToExcel(df, columns, filename):
    folder = os.path.dirname(filename)
    if folder:
        os.makedirs(folder, exist_ok=True)

    clusters = sorted(df["cluster"].unique())
    total = len(df)
    statColumns = columns + ["TotalSpent"]

    rows = []
    for c in clusters:
        dfCluster = df[df["cluster"] == c]
        row = {"Cụm": f"Cluster {c}", "Tên cụm": dfCluster["ClusterName"].iloc[0],
               "Số lượng": len(dfCluster), "Tỷ lệ (%)": round(len(dfCluster) / total * 100, 2)}
        for col in statColumns:
            row[f"{col} (TB)"] = round(dfCluster[col].mean(), 2)
            row[f"{col} (Min)"] = dfCluster[col].min()
            row[f"{col} (Max)"] = dfCluster[col].max()
        rows.append(row)
    totalRow = {"Cụm": "Tổng cộng", "Tên cụm": "", "Số lượng": total, "Tỷ lệ (%)": 100.0}
    for col in statColumns:
        totalRow[f"{col} (TB)"] = round(df[col].mean(), 2)
        totalRow[f"{col} (Min)"] = df[col].min()
        totalRow[f"{col} (Max)"] = df[col].max()
    rows.append(totalRow)
    dfStats = pd.DataFrame(rows)

    headerFont = Font(bold=True, color="FFFFFF")
    headerFill = PatternFill("solid", fgColor="305496")

    def formatSheet(ws):
        for cell in ws[1]:
            cell.font = headerFont
            cell.fill = headerFill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.freeze_panes = "A2"
        for colCells in ws.columns:
            maxLen = max(len(str(cell.value)) if cell.value is not None else 0 for cell in colCells)
            ws.column_dimensions[colCells[0].column_letter].width = min(maxLen + 3, 40)

    with pd.ExcelWriter(filename, engine="openpyxl") as writer:
        for c in clusters:
            sheetName = f"Cluster {c}"
            dfCluster = df[df["cluster"] == c]
            dfCluster.to_excel(writer, sheet_name=sheetName, index=False)
            formatSheet(writer.sheets[sheetName])
        dfStats.to_excel(writer, sheet_name="Thống kê", index=False)
        wsStats = writer.sheets["Thống kê"]
        formatSheet(wsStats)
        for cell in wsStats[wsStats.max_row]:   # tô đậm dòng tổng cộng
            cell.font = Font(bold=True)

    print("Đã xuất Excel:", os.path.abspath(filename))
    return dfStats


# ======================================================================
# 2. XUẤT HTML: 1 trang, mỗi thực nghiệm 1 tab, biểu đồ plotly tương tác
#    experiments: danh sách dict do Clustering/export_results.py chuẩn bị
# ======================================================================
def exportClustersToHtml(overview, experiments, dfLookup, filename):
    def figToHtml(fig, height=420):
        bottom = fig.layout.margin.b if fig.layout.margin.b is not None else 50
        fig.update_layout(autosize=True, margin=dict(l=50, r=20, t=60, b=bottom),
                          font=dict(family="Segoe UI, Arial", size=13))
        return fig.to_html(full_html=False, include_plotlyjs=False,
                           default_width="100%", default_height=f"{height}px",
                           config={"displaylogo": False, "responsive": True})

    def tableHtml(data, escape=True):
        return data.to_html(index=False, classes="data-table", border=0, escape=escape,
                            float_format=lambda x: f"{x:,.2f}")

    tabButtons = '<button class="tab active" data-tab="tab-overview">Tổng quan</button>'
    tabPanels = ""

    # ---------------- Tab tổng quan ----------------
    dfCompare = pd.DataFrame([{
        "Thực nghiệm": e["name"], "Thuộc tính": ", ".join(e["columns"]),
        "Khoảng elbow": f"{min(e['elbowZone'])} - {max(e['elbowZone'])}",
        "K (Kneedle)": e["k"], "WCSS": f'{e["wcss"]:.4f}', "Silhouette": f'{e["silhouette"]:.4f}'}
        for e in experiments])
    cards = "".join(f'<div class="card">{label}<b>{value}</b></div>' for label, value in overview["cards"])
    tabPanels += f"""
    <div class="panel active" id="tab-overview">
      <div class="cards">{cards}</div>
      <section id="sec-overview-compare">
        <h2>So sánh 3 thực nghiệm</h2>
        <p class="meta">K được chọn tự động bằng thuật toán Kneedle (điểm gập lớn nhất của đồ thị Elbow).
           Silhouette càng gần 1 thì các cụm càng tách biệt.</p>
        {tableHtml(dfCompare)}
      </section>
      <section id="sec-overview-features">
        <h2>Thuộc tính dùng để gom cụm</h2>
        {tableHtml(overview["features"])}
      </section>
    </div>"""

    # ---------------- Tab từng thực nghiệm ----------------
    for idx, e in enumerate(experiments, start=1):
        tabId = f"tab-exp{idx}"
        tabButtons += f'<button class="tab" data-tab="{tabId}">{html.escape(e["name"])}</button>'
        df = e["df"]
        clusters = sorted(df["cluster"].unique())
        labelOf = {c: f"Cluster {c}: {e['names'][c]}" for c in clusters}
        colorOf = {labelOf[c]: CLUSTER_COLORS[(c - 1) % len(CLUSTER_COLORS)] for c in clusters}
        dfPlot = df.copy()
        dfPlot["Cụm"] = dfPlot["cluster"].map(labelOf)
        order = [labelOf[c] for c in clusters]

        # Elbow (Kneedle)
        figElbow = go.Figure()
        figElbow.add_vrect(x0=min(e["elbowZone"]) - 0.3, x1=max(e["elbowZone"]) + 0.3,
                           fillcolor="orange", opacity=0.15, line_width=0,
                           annotation_text="Khoảng elbow", annotation_position="top left")
        figElbow.add_trace(go.Scatter(x=e["kRange"], y=e["inertia"], mode="lines+markers",
                                      name="WCSS", line=dict(color="#1d3557")))
        figElbow.add_trace(go.Scatter(x=[e["k"]], y=[e["inertia"][e["k"] - 1]], mode="markers+text",
                                      marker=dict(size=16, color="red"), text=[f"K = {e['k']}"],
                                      textposition="top right", name="Điểm gập lớn nhất"))
        figElbow.update_layout(title="Elbow Method + Kneedle", xaxis_title="Số cụm K",
                               yaxis_title="WCSS (inertia)", xaxis=dict(dtick=1), showlegend=False)

        # Tỷ lệ khách theo cụm
        dfCount = dfPlot.groupby("Cụm").size().reindex(order).reset_index(name="Số khách hàng")
        figPie = px.pie(dfCount, names="Cụm", values="Số khách hàng", hole=0.45,
                        title="Tỷ lệ khách hàng theo cụm", color="Cụm", color_discrete_map=colorOf,
                        category_orders={"Cụm": order})
        figPie.update_traces(textinfo="percent+value", sort=False)
        figPie.update_layout(legend=dict(orientation="h", y=-0.15))

        # Biểu đồ cụm (2D hoặc 3D)
        cols = e["columns"]
        hover = ["CustomerId", "FullName", "Country"] + cols
        if len(cols) == 2:
            figScatter = px.scatter(dfPlot, x=cols[0], y=cols[1], color="Cụm", hover_data=hover,
                                    color_discrete_map=colorOf, category_orders={"Cụm": order},
                                    title="Các cụm khách hàng - " + " X ".join(cols))
            figScatter.update_traces(marker=dict(size=12, line=dict(width=1, color="white")))
            scatterHeight = 520
        else:
            figScatter = px.scatter_3d(dfPlot, x=cols[0], y=cols[1], z=cols[2], color="Cụm",
                                       hover_data=hover, color_discrete_map=colorOf,
                                       category_orders={"Cụm": order},
                                       title="Các cụm khách hàng (3D) - " + " X ".join(cols))
            figScatter.update_traces(marker=dict(size=6))
            figScatter.update_layout(scene=dict(aspectmode="cube"))
            scatterHeight = 720
        figScatter.update_layout(legend=dict(orientation="h", yanchor="top", y=-0.18, title_text=""),
                                 margin=dict(b=120))

        # Bảng tóm tắt cụm
        dfSummary = e["summary"]

        # Từng cụm: top thể loại (biểu đồ) + top album (bảng) + danh sách khách
        clusterSections = ""
        for c in clusters:
            dfC = df[df["cluster"] == c]
            color = CLUSTER_COLORS[(c - 1) % len(CLUSTER_COLORS)]
            dfGenre = e["topGenres"][c]
            figGenre = px.bar(dfGenre, x="Tỷ lệ lượt mua (%)", y="Genre", orientation="h", text="Tỷ lệ lượt mua (%)",
                              title=f"Cluster {c}: 5 thể loại được mua nhiều nhất")
            figGenre.update_traces(marker_color=color, texttemplate="%{text:.1f}%")
            figGenre.update_layout(yaxis=dict(autorange="reversed", title=""), xaxis_title="% lượt mua của cụm")
            meta = " · ".join(f"{col} TB <b>{dfC[col].mean():.2f}</b>" for col in cols)
            clusterSections += f"""
            <section class="cluster" id="sec-exp{idx}-cluster{c}" style="border-top-color:{color}">
              <h3><span class="dot" style="background:{color}"></span>Cluster {c}: {html.escape(e['names'][c])}
                  <small>{len(dfC)} khách hàng ({len(dfC) / len(df) * 100:.1f}%)</small></h3>
              <p class="meta">{meta} · Chi tiêu TB <b>{dfC['TotalSpent'].mean():.2f} USD</b></p>
              <div class="grid">
                <div class="chart">{figToHtml(figGenre, 320)}</div>
                <div><h4>5 album được nhiều khách trong cụm mua nhất</h4>{tableHtml(e['topAlbums'][c])}</div>
              </div>
              <details><summary>Xem danh sách {len(dfC)} khách hàng của Cluster {c}</summary>
                {tableHtml(dfC[['CustomerId', 'FullName', 'Country', 'City'] + cols + ['TotalSpent']])}</details>
            </section>"""

        expCards = (f'<div class="card">Số cụm K<b>{e["k"]}</b></div>'
                    f'<div class="card">Khoảng elbow<b>{min(e["elbowZone"])} - {max(e["elbowZone"])}</b></div>'
                    f'<div class="card">Silhouette<b>{e["silhouette"]:.3f}</b></div>'
                    f'<div class="card">WCSS<b>{e["wcss"]:.3f}</b></div>')
        tabPanels += f"""
    <div class="panel" id="{tabId}">
      <div class="cards">{expCards}</div>
      <section id="sec-exp{idx}-elbow">
        <h2>Chọn số cụm và tỷ lệ khách hàng</h2>
        <p class="meta">Thuộc tính: <b>{', '.join(cols)}</b> (chuẩn hóa Min-Max trước khi gom cụm).</p>
        <div class="grid"><div class="chart">{figToHtml(figElbow)}</div><div class="chart">{figToHtml(figPie)}</div></div>
      </section>
      <section id="sec-exp{idx}-scatter">
        <h2>Biểu đồ các cụm</h2>
        <div class="chart">{figToHtml(figScatter, scatterHeight)}</div>
      </section>
      <section id="sec-exp{idx}-summary">
        <h2>Thống kê theo cụm</h2>
        {tableHtml(dfSummary)}
      </section>
      {clusterSections}
    </div>"""

    # ---------------- Tab tra cứu & gợi ý ----------------
    tabButtons += '<button class="tab" data-tab="tab-lookup">Tra cứu &amp; gợi ý album</button>'
    tabPanels += f"""
    <div class="panel" id="tab-lookup">
      <section id="sec-lookup">
        <h2>Cụm của từng khách hàng và 5 album gợi ý</h2>
        <p class="meta">Gợi ý = album được nhiều khách cùng cụm (mô hình ALBUM) mua nhất mà khách này chưa mua.</p>
        <input id="search" type="text" placeholder="Nhập tên, quốc gia hoặc CustomerId để lọc...">
        {tableHtml(dfLookup, escape=False).replace('class="dataframe data-table"', 'class="dataframe data-table" id="lookup-table"')}
      </section>
    </div>"""

    page = f"""<!DOCTYPE html>
<html lang="vi">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Chinook - Gom cụm khách hàng</title>
<!-- plotly.js nhúng 1 lần ở đầu trang: file xem được cả khi offline -->
<script>{get_plotlyjs()}</script>
<style>
  body {{ font-family: "Segoe UI", Arial, sans-serif; margin: 0; background: #f4f6fb; color: #222; }}
  header {{ background: #1d3557; color: #fff; padding: 26px 40px 0; }}
  header h1 {{ margin: 0 0 6px; font-size: 26px; }}
  header p {{ margin: 0 0 16px; color: #d6e2f0; }}
  .tabs {{ display: flex; gap: 4px; flex-wrap: wrap; }}
  .tab {{ background: #2b4a75; color: #d6e2f0; border: 0; padding: 10px 18px; font-size: 15px;
          border-radius: 8px 8px 0 0; cursor: pointer; }}
  .tab.active {{ background: #f4f6fb; color: #1d3557; font-weight: 600; }}
  main {{ max-width: 1240px; margin: 0 auto; padding: 24px; }}
  .panel {{ display: none; }}
  .panel.active {{ display: block; }}
  .cards {{ display: flex; gap: 16px; flex-wrap: wrap; margin-bottom: 20px; }}
  .card {{ background: #fff; border-radius: 10px; padding: 14px 22px; flex: 1; min-width: 160px;
          box-shadow: 0 1px 4px rgba(0,0,0,.08); color: #555; }}
  .card b {{ display: block; font-size: 26px; color: #1d3557; margin-top: 4px; }}
  section {{ background: #fff; border-radius: 10px; padding: 20px; margin-bottom: 22px;
            box-shadow: 0 1px 4px rgba(0,0,0,.08); }}
  section.cluster {{ border-top: 6px solid; }}
  h2 {{ margin: 0 0 10px; color: #1d3557; }}
  h3 {{ margin: 0 0 6px; }}
  h3 small {{ font-weight: normal; color: #666; font-size: 14px; margin-left: 8px; }}
  h4 {{ margin: 8px 0; color: #333; }}
  .dot {{ display: inline-block; width: 14px; height: 14px; border-radius: 50%; margin-right: 8px; }}
  .meta {{ color: #444; line-height: 1.7; margin: 0 0 10px; }}
  .grid {{ display: grid; grid-template-columns: repeat(2, minmax(0, 1fr)); gap: 16px; align-items: start; }}
  .chart {{ min-width: 0; }}
  .data-table {{ border-collapse: collapse; width: 100%; font-size: 13.5px; margin-top: 8px; }}
  .data-table thead tr {{ text-align: center !important; }}
  .data-table th {{ background: #1d3557; color: #fff; padding: 8px; position: sticky; top: 0; text-align: center; }}
  .data-table td:first-child {{ white-space: nowrap; }}
  .data-table td {{ padding: 6px 8px; border-bottom: 1px solid #e5e5e5; text-align: center; }}
  .data-table tr:nth-child(even) td {{ background: #f8f9fc; }}
  details {{ margin-top: 12px; }}
  details table {{ max-height: 380px; overflow: auto; display: block; }}
  summary {{ cursor: pointer; color: #1d3557; font-weight: 600; }}
  #search {{ width: 100%; box-sizing: border-box; padding: 10px 12px; font-size: 15px;
             border: 1px solid #c5cbd6; border-radius: 8px; margin: 6px 0 10px; }}
  #lookup-table {{ display: block; max-height: 640px; overflow: auto; }}
  #lookup-table td:last-child {{ text-align: left; font-size: 13px; }}
  footer {{ text-align: center; color: #777; padding: 10px 0 30px; font-size: 13px; }}
  @media (max-width: 900px) {{ .grid {{ grid-template-columns: 1fr; }} main {{ padding: 12px; }} }}
</style>
</head>
<body>
<header>
  <h1>Chinook Music Store - Gom cụm khách hàng theo mức độ quan tâm Track &amp; Album</h1>
  <p>K-Means (k-means++, n_init = 10, random_state = 42) · chọn K bằng Elbow + Kneedle · đánh giá bằng Silhouette</p>
  <nav class="tabs">{tabButtons}</nav>
</header>
<main>
{tabPanels}
</main>
<footer>{html.escape(overview["footer"])}</footer>
<script>
  // Chuyển tab; vẽ lại biểu đồ plotly của tab vừa mở cho đúng kích thước khung
  document.querySelectorAll(".tab").forEach(function (btn) {{
    btn.addEventListener("click", function () {{
      document.querySelectorAll(".tab").forEach(function (b) {{ b.classList.remove("active"); }});
      document.querySelectorAll(".panel").forEach(function (p) {{ p.classList.remove("active"); }});
      btn.classList.add("active");
      var panel = document.getElementById(btn.dataset.tab);
      panel.classList.add("active");
      panel.querySelectorAll(".js-plotly-plot").forEach(function (d) {{ Plotly.Plots.resize(d); }});
    }});
  }});
  // Lọc bảng tra cứu theo từ khóa
  var search = document.getElementById("search");
  if (search) {{
    search.addEventListener("input", function () {{
      var key = search.value.toLowerCase();
      document.querySelectorAll("#lookup-table tbody tr").forEach(function (tr) {{
        tr.style.display = tr.innerText.toLowerCase().indexOf(key) >= 0 ? "" : "none";
      }});
    }});
  }}
  window.addEventListener("load", function () {{
    document.querySelectorAll(".panel.active .js-plotly-plot").forEach(function (d) {{ Plotly.Plots.resize(d); }});
  }});
</script>
</body>
</html>"""

    folder = os.path.dirname(filename)
    if folder:
        os.makedirs(folder, exist_ok=True)
    with open(filename, "w", encoding="utf-8") as f:
        f.write(page)
    print("Đã xuất HTML:", os.path.abspath(filename))
